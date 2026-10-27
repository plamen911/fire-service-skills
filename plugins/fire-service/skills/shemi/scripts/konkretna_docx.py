#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
konkretna_docx.py — КОНКРЕТНАТА СХЕМА С РЕШЕНИЯ ЗА ГАСЕНЕ като .docx за печат (чл. 7, ал. 4 от Указанията
Iз-2749/2010) – два листа A4, легнали, по чл. 9, ал. 6 в общ джоб: на лицевата страна графичната част
(Приложение № 3а – листът с рамката, полето за номера и антетката), на гърба – табличната (Приложение № 3б).

    python3 scripts/konkretna_docx.py konkretna.json -o "/mnt/user-data/outputs/Konkretna_shema_<Име>.docx"

konkretna.json е изходът на scripts/konkretna.py (входът на shema.py, с "list") и "tablitsa_3b":
{"redove": [{"kategoriya": "Ф1.1", "ogneustoychivost": "II", "natovarvane_vid": "текстил", "natovarvane_kg_m2": "≈ 86",
             "pgi_vid": "няма", "vl_m_min": "0,4", "plamna_c": "", "gasitelni_neobhodimi": "вода",
             "gasitelni_zabraneni": "вода върху ел. уредби под напрежение", "zashtita": "ВДА, ЛПС",
             "rvm_vid": "няма", "drugi": "…"}], "prazni_redove": 4}
konkretna.py попълва каквото знае (категорията, огнеустойчивостта, линейната скорост от разчета); останалото –
от данните за обекта. Празна клетка остава празна – не се измисля.

До .docx-а се записват графичната част като картинка (…_grafichna.png) и като .dxf за AutoCAD/CorelDRAW.
Отпечатва JSON {"output", "grafichna", "dxf", "mashtab", "preduprezhdeniya"}.
"""
import argparse
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "_shared", "scripts"))

import shema  # noqa: E402
from prilozhenia import FONT, number_box, page, para, table_3b  # noqa: E402
from situacionna import pad_to_page  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

try:
    from docx_common import report, typo_all  # noqa: E402
except ImportError:  # pragma: no cover – the development tree
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(HERE))), "shared", "scripts"))
    from docx_common import report, typo_all  # noqa: E402


def bad(msg):
    sys.stderr.write("BAD_INPUT: " + msg + "\n")
    sys.exit(2)


def gaps_3b(tab):
    """What чл. 7, ал. 4, т. 2 asks for and the table does not have."""
    rows = tab.get("redove") or []
    out = []
    for key, what, letter in (("kategoriya", "категорията по пожарна опасност", "а"),
                              ("ogneustoychivost", "степента на огнеустойчивост", "б"),
                              ("natovarvane_vid", "основното горимо натоварване", "в"),
                              ("vl_m_min", "линейната скорост на разпространение", "д"),
                              ("gasitelni_neobhodimi", "необходимите гасителни вещества", "ж"),
                              ("zashtita", "вида защита на служителите", "з")):
        if not any(str(r.get(key) or "").strip() for r in rows):
            out.append(f"таблична част (Приложение № 3б): няма {what} (чл. 7, ал. 4, т. 2, б. „{letter}“)")
    return out


def razchet_module():
    """razchet_docx of the skill sili-i-sredstva – it sits beside shemi in the plugin."""
    path = os.path.join(os.path.dirname(os.path.dirname(HERE)), "sili-i-sredstva", "scripts")
    if path not in sys.path:
        sys.path.insert(0, path)
    import razchet_docx
    return razchet_docx


def back_with_calculation(doc, tab, sili, obekt_line="", zashtita=None, naslov=None, extra_cm=0.0):
    """The table of Приложение № 3б and under it, on the same back page, the calculation (Приложение № 5) – one
    pocket per scheme (чл. 9, ал. 6). Returns the warnings."""
    from prilozhenia import COLUMNS_3B, WIDTHS_3B, table_height_cm
    rows = list(tab.get("redove") or [])
    tab = dict(tab, prazni_redove=len(rows))           # no empty rows – the room goes to the calculation
    table_3b(doc, tab)
    used = extra_cm + 1.5 + 1.1 + table_height_cm(COLUMNS_3B, WIDTHS_3B, rows, max(len(rows), tab["prazni_redove"]))
    return razchet_module().page(doc, sili, naslov, obekt_line, None, under_cm=used, zashtita=zashtita)


def build(d, out, sili=None, zashtita=None):
    if not isinstance(d, dict) or not isinstance(d.get("list"), dict):
        bad("очаква се JSON-ът на konkretna.py с \"list\"")
    tab = d.get("tablitsa_3b") or {}
    if not isinstance(tab, dict):
        bad("tablitsa_3b: очаква се обект")
    lst = d["list"]
    base = os.path.splitext(out)[0]
    png, dxf = base + "_grafichna.png", base + ".dxf"
    if not str(lst.get("fail") or "").strip():
        lst["fail"] = os.path.basename(dxf)
    _, _, warnings, extra = shema.render(d, png, out_dxf=dxf)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    page(doc.sections[0])
    p = para(doc)                                      # 1. the graphic part – the whole framed sheet (Прил. № 3а)
    p.add_run().add_picture(png, width=Cm(26.6))
    sec = doc.add_section(WD_SECTION.NEW_PAGE)        # 2. the table part – on the back of the sheet (Прил. № 3б)
    page(sec)
    number_box(doc, lst.get("nomer"))
    para(doc, str(lst.get("naslov") or "Конкретна схема") + " – таблична част", 12, WD_ALIGN_PARAGRAPH.CENTER)
    if lst.get("obekt"):
        para(doc, str(lst["obekt"]), 12, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    calc_warnings = []
    try:
        if sili:                                         # the calculation on the same back page, under the table
            calc_warnings = back_with_calculation(doc, tab, sili, zashtita=zashtita, naslov=lst.get("naslov"))
            n = len(tab.get("redove") or [])
        else:
            n = table_3b(doc, tab)
    except ValueError as e:
        bad(str(e))
    typo_all(doc)
    doc.save(out)
    report(out, letterhead=False)
    pad_to_page(png)
    # what the graphic part (чл. 7, ал. 4, т. 1) is to show and the scheme lacks – the scheme's own warnings that
    #   speak of the situational scheme (streets, mains …) do not apply here
    w = [x for x in warnings if "б. „" not in x]
    return {"output": out, "grafichna": png, "dxf": extra.get("dxf"), "mashtab": extra.get("mashtab"), "redove": n,
            "preduprezhdeniya": w + gaps_3b(tab) + calc_warnings}


def main():
    ap = argparse.ArgumentParser(description="Конкретна схема с решения за гасене (Приложения № 3а и № 3б)")
    ap.add_argument("data", help="JSON-ът на konkretna.py (с \"list\" и \"tablitsa_3b\")")
    ap.add_argument("-o", "--out", required=True, help="изходен .docx")
    ap.add_argument("--sili", help="JSON-ът на sili.py – разчетът отива на гърба, под таблицата № 3б")
    ap.add_argument("--zashtita", help="какво се защитава – за реда на разхода за защита в разчета")
    a = ap.parse_args()
    try:
        d = json.load(open(a.data, encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        bad(f"не мога да прочета {a.data}: {e}")
    os.makedirs(os.path.dirname(os.path.abspath(a.out)) or ".", exist_ok=True)
    sili = None
    if a.sili:
        try:
            sili = json.load(open(a.sili, encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as e:
            bad(f"не мога да прочета {a.sili}: {e}")
    print(json.dumps(build(d, a.out, sili, a.zashtita), ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
