#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
situacionna.py — СИТУАЦИОННА СХЕМА на обект по чл. 5, ал. 2 и чл. 7, ал. 3 от Указанията за разработване на
планове за пожарогасене (рег. № Iз-2749/18.11.2010 г. на МВР) като .docx за печат – два листа A4, легнали:

  1. графична част (Приложение № 2а) – листът с рамката, полето за номера и антетката; чертае го
     scripts/shema.py (същият JSON, с "list"), знаците са от библиотеката scripts/znatsi.py;
  2. таблична част (Приложение № 2б) – таблицата като таблица на Word (попълва се и на ръка), полето за
     номера горе вдясно (чл. 9, ал. 3) и легендата с условните знаци (чл. 9, ал. 10).

По чл. 9, ал. 6 графичната част е на лицевата страна, а табличната – на гърба ѝ: печата се двустранно.

    python3 scripts/situacionna.py situacionna.json -o "/mnt/user-data/outputs/Situacionna_shema_<Обект>.docx"

situacionna.json е JSON-ът на схемата (входът – в началото на scripts/shema.py), с "list" и още:
{
  ...полетата на схемата: "sgradi" (с "etazhnost", "kategoriya", "ogneustoychivost"), "ulitsi", "vodoprovod",
     "vodoiztochnitsi", "gazoprovod", "gaz_armatura", "bez_parkirane", "razstoyaniya", "obekti" („Ел.“) ...,
  "list": {"nomer": "7", "obekt": "…", "sobstvenik": "…", "data": "24.10.2026", ...},
                                               // "izgotvil" по подразбиране – потребителят по личния ключ;
                                               // "fail" – името на .dxf файла на графичната част
                                               //   (с конектора – името от whoami); "data" – днешната дата
  "kadastar": "56722.651.111",                // по желание – иначе от "kais.imot": землището (ЕКАТТЕ) определя службата
  "obekt_gps": "43.4125702,24.6504731",       // точката на обекта или линк към нея – тогава разстоянието и
                                               //   времето се смятат сами от службата по заповедта (marshrut.py:
                                               //   по пътищата, 1 km/min и 1 min излизане); "sluzhba": "…" – от друга служба
  "tablitsa": {
    "razstoyanie_km": 3.2,                     // от службата до обекта (иначе – от "obekt_gps")
    "vreme_min": 7,                            // времето за реагиране (иначе – излизане + движение от "obekt_gps")
    "redove": [{"natovarvane_vid": "мебели, текстил", "natovarvane_kg_m2": 25,
                "pgi_vid": "", "pgi_kolichestvo": "",            // пожарогасителни инсталации: вид и к-во гас. в-во
                "gasitelni_neobhodimi": "вода", "gasitelni_zabraneni": "",
                "rvm_vid": "", "rvm_kolichestvo": "",            // рискови вещества и материали
                "gazove_vid": "", "gazove_kolichestvo": "",
                "zashtita": "ДА"}],                              // вид защита на служителите
    "prazni_redove": 6                         // общо редове в таблицата (празните се попълват на ръка)
  },
  "legenda": true                              // легендата под таблицата (по подразбиране – да)
}

Празна клетка остава празна – стойност, която не е дадена, не се измисля.
Типографията следва _shared/conventions.md: „№ 4“, тире „–“, без получер шрифт; без шапка (бланката е
листът по Приложение № 2а).

Отпечатва JSON {"output", "grafichna", "mashtab", "redove", "preduprezhdeniya"}. Предупрежденията казват
какво от чл. 7, ал. 3 липсва в графичната и в табличната част.

Грешки (на stderr, код ≠ 0):
    BAD_INPUT: ...   — неверен JSON, липсва "sgradi" или "list"
"""
import argparse
import datetime
import json
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "_shared", "scripts"))

import shema  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

try:
    from docx_common import report, typo, typo_all  # noqa: E402
except ImportError:  # pragma: no cover – the skill without _shared (development tree)
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(HERE))), "shared", "scripts"))
    from docx_common import report, typo, typo_all  # noqa: E402

from prilozhenia import (FONT, MARGINS, borders, cell_text, fixed_width, fmt, number_box, page,  # noqa: E402,F401
                         para, row_height, table_2b as _table_2b)

def bad(msg):
    sys.stderr.write("BAD_INPUT: " + msg + "\n")
    sys.exit(2)


def table_2b(doc, tab):
    try:
        return _table_2b(doc, tab)
    except ValueError as e:
        bad(str(e))


def table_gaps(tab):
    gaps = []
    if tab.get("razstoyanie_km") in (None, ""):
        gaps.append("таблична част: няма разстояние до обекта – \"razstoyanie_km\" (чл. 7, ал. 3, т. 2, б. „а“)")
    if tab.get("vreme_min") in (None, ""):
        gaps.append("таблична част: няма време за реагиране – \"vreme_min\" (б. „б“)")
    rows = tab.get("redove") or []
    for key, what, letter in (("natovarvane_vid", "горимото натоварване", "в"),
                              ("gasitelni_neobhodimi", "необходимите гасителни вещества", "д"),
                              ("zashtita", "вида защита на служителите", "з")):
        if not any(str(r.get(key) or "").strip() for r in rows):
            gaps.append(f"таблична част: няма {what} (б. „{letter}“)")
    return gaps


def own_name():
    """Име и фамилия на потребителя по личния му ключ (списъкът на служителите – _shared/profile.md);
    None, когато ключ няма (тогава "izgotvil" се подава от whoami на конектора)."""
    script = os.path.join(os.path.dirname(HERE), "_shared", "scripts", "sluzhiteli.py")
    if not os.path.exists(script):
        script = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(HERE))), "shared", "scripts", "sluzhiteli.py")
    try:
        r = subprocess.run([sys.executable, script, "--profil"], capture_output=True, text=True, timeout=90)
        return (json.loads(r.stdout).get("profil") or {}).get("ime") if r.returncode == 0 else None
    except (OSError, ValueError, subprocess.TimeoutExpired):
        return None


def kadastar_id(d):
    """The cadastral identifier of the parcel (from kais.py): "kadastar", "kais.imot" or the label of "imot"."""
    for v in (d.get("kadastar"), (d.get("kais") or {}).get("imot"), (d.get("imot") or {}).get("etiket")):
        m = re.search(r"\d{5}\.\d+\.\d+", str(v or ""))
        if m:
            return m.group(0)
    return None


def route_to(target, ot=None, kadastar=None):
    """Distance and travel time from the responsible service (marshrut.py); None when it cannot be computed."""
    script = os.path.join(os.path.dirname(HERE), "_shared", "scripts", "marshrut.py")
    if not os.path.exists(script):
        script = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(HERE))), "shared", "scripts", "marshrut.py")
    cmd = [sys.executable, script, "--link" if "://" in target else "--do", target] + (["--ot", ot] if ot else [])
    if kadastar:                       # the land of the parcel (ЕКАТТЕ) decides the settlement, not the nearest centre
        cmd += ["--kadastar", kadastar]
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        return json.loads(r.stdout) if r.returncode == 0 else None
    except (OSError, ValueError, subprocess.TimeoutExpired):
        return None


def build(d, out):
    if not isinstance(d, dict):
        bad("очаква се JSON обект")
    if not isinstance(d.get("list"), dict):
        bad("липсва \"list\" – полетата на антетката (поне \"obekt\")")
    tab = d.get("tablitsa") or {}
    if not isinstance(tab, dict):
        bad("tablitsa: очаква се обект")
    lst = d["list"]
    gaps = []
    if not str(lst.get("izgotvil") or "").strip():         # „Изготвил“ – потребителят по ключа
        name = own_name()
        if name:
            lst["izgotvil"] = name
        else:
            gaps.append("„Изготвил“ е празно – подай \"list.izgotvil\" (името от whoami на конектора)")
    target = str(d.get("obekt_gps") or "").strip()
    marshrut = None
    if target and (tab.get("razstoyanie_km") in (None, "") or tab.get("vreme_min") in (None, "")):
        marshrut = route_to(target, d.get("sluzhba"), kadastar_id(d))   # the table's distance and time – from the service
        if marshrut:
            if tab.get("razstoyanie_km") in (None, ""):
                tab["razstoyanie_km"] = round(marshrut["razstoyanie_km"], 1)
            if tab.get("vreme_min") in (None, ""):
                tab["vreme_min"] = marshrut["vreme_min"]
        else:
            gaps.append("разстоянието и времето не се изчислиха (marshrut.py) – пусни го отделно и ги подай в \"tablitsa\"")
    if not str(lst.get("data") or "").strip():
        lst["data"] = datetime.date.today().strftime("%d.%m.%Y")
    base = os.path.splitext(out)[0]
    png = base + "_grafichna.png"
    dxf = base + ".dxf"                                # the graphic part for AutoCAD / CorelDRAW
    if not str(lst.get("fail") or "").strip():         # „Име на файла“ – the drawing's own file
        lst["fail"] = os.path.basename(dxf)
    _, _, warnings, extra = shema.render(d, png, out_dxf=dxf)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    page(doc.sections[0])
    # 1. the graphic part – the whole framed sheet
    p = para(doc)
    p.add_run().add_picture(png, width=Cm(26.6))
    # 2. the table part – on a new landscape page (the back of the sheet)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page(sec)
    number_box(doc, lst.get("nomer"))
    para(doc, "Ситуационна схема – таблична част", 12, WD_ALIGN_PARAGRAPH.CENTER)
    if lst.get("obekt"):
        para(doc, str(lst["obekt"]), 12, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    n = table_2b(doc, tab)
    if d.get("legenda", True):
        leg = base + "_legenda.png"
        if shema.render_legend(d, leg, width_cm=20.0):
            para(doc, "", after=4)
            q = para(doc)
            q.add_run().add_picture(leg, width=Cm(20.0))
    typo_all(doc)
    doc.save(out)
    report(out, letterhead=False)
    pad_to_page(png)                                   # the .png on its own: the whole A4 page, with its margins
    return {"output": out, "grafichna": png, "dxf": extra.get("dxf"), "mashtab": extra.get("mashtab"), "redove": n,
            "izgotvil": lst.get("izgotvil") or None, "marshrut": marshrut,
            "preduprezhdeniya": warnings + gaps + table_gaps(tab)}


def pad_to_page(png, dpi=300):
    """The framed sheet as a whole A4 page: white margins round the frame (left 2 cm for binding, 1 cm on the
    other sides – the margins of the .docx), so the picture printed or viewed on its own has room round the
    frame. Called after the picture went into the .docx, where the page margins already give that room."""
    try:
        from PIL import Image
    except ImportError:                                # pragma: no cover – Pillow comes with matplotlib
        return
    mm = dpi / 25.4
    im = Image.open(png)
    l, r, t, b = (round(MARGINS[k] * 10 * mm) for k in ("left", "right", "top", "bottom"))
    page_ = Image.new("RGB", (im.width + l + r, im.height + t + b), "white")
    page_.paste(im.convert("RGB"), (l, t))
    page_.save(png, dpi=(dpi, dpi))


def main():
    ap = argparse.ArgumentParser(description="Ситуационна схема на обект (Приложения № 2а и № 2б към Iз-2749)")
    ap.add_argument("data", help="JSON на схемата с \"list\" и \"tablitsa\"")
    ap.add_argument("-o", "--out", required=True, help="изходен .docx")
    a = ap.parse_args()
    try:
        with open(a.data, encoding="utf-8") as f:
            d = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        bad(f"не мога да прочета {a.data}: {e}")
    os.makedirs(os.path.dirname(os.path.abspath(a.out)) or ".", exist_ok=True)
    print(json.dumps(build(d, a.out), ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
