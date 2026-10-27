#!/usr/bin/env python3
"""build_plan.py on a made-up object: the plan and the sheets on their own, the gaps it reports."""
import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
SKILLS = os.path.dirname(SKILL)
FIX = os.path.join(SKILLS, "shemi", "tests", "fixtures", "situacionna.json")


def run(args, cwd=None):
    p = subprocess.run([sys.executable, *args], capture_output=True, text=True, cwd=cwd)
    assert p.returncode == 0, p.stderr[-1500:]
    return json.loads(p.stdout)


def main():
    from docx import Document
    with tempfile.TemporaryDirectory() as d:
        s = run([os.path.join(SKILLS, "sili-i-sredstva", "scripts", "sili.py"), "tvardi", "--forma", "pravoagalnik",
                 "--fp", "40", "--a", "8", "--b", "5", "--vl-red", "1", "--i-red", "6.3", "--ataka", "front1",
                 "--struinik", "В", "--q-pa", "40"])
        sit = json.load(open(FIX, encoding="utf-8"))
        k = json.loads(json.dumps(sit))
        k["list"]["naslov"] = "Конкретна схема за гасене на пожар в склад на 1 ет."
        k["tablitsa_3b"] = {"redove": [{"kategoriya": "Ф5.2", "ogneustoychivost": "III", "natovarvane_vid": "дърво",
                                        "vl_m_min": "1,5", "gasitelni_neobhodimi": "вода", "zashtita": "ВДА"}]}
        plan = {"fail": "Plan_pozharogasene_Test", "godina": 2026, "obekt": "Склад „Тест“", "adres": "с. Тестово",
                "utvarzhdava": {"lines": ["ДИРЕКТОР НА РДПБЗН – ПЛЕВЕН"], "name": "Име Фамилия"},
                "danni": [["Разположение на сградите.", "Една сграда."], ["Подстъпи.", "Асфалт."]],
                "situacionna": sit,
                "varianti": [{"ime": "пожар в склада", "fail": "Konkretna_shema_Sklad_1et", "konkretna": k, "sili": s,
                              "pomeshtenie": "Склад – 8 × 5 m",
                              "reshenie": [["Място на пожара.", "Складът."], ["Схема на подаване.", "I ПА – пред входа."]]}],
                "izgotvil": {"lines": ["инспектор"], "name": "Тест Тестов", "data": "27.10.2026"}}
        src = os.path.join(d, "plan.json")
        json.dump(plan, open(src, "w", encoding="utf-8"), ensure_ascii=False)
        r = run([os.path.join(SKILL, "scripts", "build_plan.py"), src, "-d", d])
        names = {os.path.basename(f) for f in r["failove"]}
        assert {"Situacionna_shema_Test.docx", "Konkretna_shema_Sklad_1et.docx", "Konkretna_shema_Sklad_1et.dxf",
                "Razchet_Konkretna_shema_Sklad_1et.docx"} <= names, names
        assert os.path.basename(r["plan"]) == "Plan_pozharogasene_Test.docx"
        doc = Document(r["plan"])
        text = "\n".join(p.text for p in doc.paragraphs)
        cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        assert "Брой струйници (общо)" in text and "Nстр" in text, "the written calculation"
        for want in ("П Л А Н", "1. ДАННИ ЗА ОБЕКТА", "ВАРИАНТ № 1", "По разчета (Приложение № 5)", "Приложение № 3б",
                     "Приложение № 4"):
            assert want in text, want
        for want in ("Ф5.2", "Произшествие или занятие"):
            assert want in cells, (want, [l for l in cells.splitlines() if "Площ" in l or "(1" in l][:10])
        assert any("т. 7" in w for w in r["preduprezhdeniya"]), r["preduprezhdeniya"]    # „вероятни опасности“ missing
        assert len(doc.sections) >= 8, len(doc.sections)
        from docx.oxml.ns import qn
        cols = [(s._sectPr.find(qn("w:cols")).get(qn("w:num")) if s._sectPr.find(qn("w:cols")) is not None else "1")
                for s in doc.sections]
        assert cols.count("2") == 1 and cols[-1] == "1", cols      # only the calculation is in two columns
        print("  ok  plan, the sheets on their own and the gaps")
    print("OK plan-pozharogasene")


if __name__ == "__main__":
    main()
