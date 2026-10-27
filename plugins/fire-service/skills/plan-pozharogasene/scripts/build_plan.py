#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_plan.py — ПЛАНЪТ ЗА ПОЖАРОГАСЕНЕ на обект по Указанията рег. № Iз-2749/18.11.2010 г. (чл. 5 – 9), сглобен от
това, което чертаят и смятат другите скилове, като .docx на бланката на РДПБЗН – Плевен, и до него – листовете
поотделно за печат:

    <fail>.docx                                    – планът: „УТВЪРЖДАВАМ“, данните за обекта (чл. 6, ал. 1),
                                                     ситуационната схема (Приложения № 2а и № 2б), за всеки вариант
                                                     решението (чл. 6, ал. 2), конкретната схема (№ 3а и № 3б) и
                                                     разчетът (№ 5), накрая бланката (№ 4)
    Situacionna_shema_<обект>.docx / .dxf / _grafichna.png
    Konkretna_shema_<вариант>.docx / .dxf / _grafichna.png    (за всеки вариант)
    Razchet_Konkretna_shema_<вариант>.docx                     (за всеки вариант)

    python3 scripts/build_plan.py plan.json -d "/mnt/user-data/outputs"

plan.json (пътищата – спрямо папката на plan.json):
{
 "fail": "Plan_pozharogasene_<Обект>",          # името на плана; „<Обект>“ дава и името на ситуационната схема
 "godina": 2026,
 "obekt": "ДГ „Слънце“", "adres": "с. Тестово, ул. …",
 "identifikator": "ПИ 00000.1.2, сграда 00000.1.2.1",
 "utvarzhdava": {"lines": [...], "name": "…"}, # по подразбиране – директорът на РДПБЗН – Плевен от списъка
 "saglasuvano": {"lines": ["Директор на …"], "name": ""},
 "danni": [["Разположение на сградите.", "текст"], …],        # чл. 6, ал. 1, т. 1 – 13
 "situacionna": "situacionna.json",              # входът на shemi/scripts/situacionna.py (с "list" и "tablitsa")
                                                 # – или самият JSON (и "konkretna", "sili" – също); такъв, цял,
                                                 # се пази в регистъра като plan_za_pozharogasene.json
 "varianti": [{"ime": "пожар в хора на 1 етаж", "fail": "Konkretna_shema_Hor_1et",
               "konkretna": "k_hor.json",        # изходът на shemi/scripts/konkretna.py (с "tablitsa_3b")
               "sili": "s_hor.json",             # изходът на sili-i-sredstva/scripts/sili.py
               "pomeshtenie": "Хор, 1 етаж – 6,9 × 10,5 m",
               "reshenie": [["Място на пожара.", "…"], …]}],          # чл. 6, ал. 2; „Сили и средства.“ – от разчета
 "prilozhenie_4": {"redove": [], "broi": 14},
 "izgotvil": {"lines": ["инспектор …"], "name": "…", "data": "27.10.2026"}
}
Отпечатва JSON {"plan", "failove", "preduprezhdeniya"}; предупрежденията казват какво от Указанията липсва.
"""
import argparse
import datetime
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SKILLS = os.path.dirname(os.path.dirname(HERE))
for p in (os.path.join(SKILLS, "plan-konspekti", "scripts"), os.path.join(SKILLS, "shemi", "scripts"),
          os.path.join(SKILLS, "sili-i-sredstva", "scripts")):
    sys.path.insert(0, p)

import blanka as B  # noqa: E402  (the letterhead and the approval – one look for the plugin)
import konkretna_docx  # noqa: E402
import razchet_docx  # noqa: E402
import situacionna  # noqa: E402
from prilozhenia import MARGINS, number_box, page, table_2b, table_3b, table_4  # noqa: E402

from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402
from docx.shared import Cm  # noqa: E402

# what чл. 6, ал. 1 asks to be studied – a heading of "danni" is matched by these words
DANNI_6_1 = [("разположение", "т. 1 – разположението на сградите"),
             ("предназначение", "т. 2 – предназначението и разпределението"),
             ("вещества", "т. 3 – съхраняваните вещества и материали"),
             ("комуникац", "т. 4 – комуникациите (вода, газ, ел., отопление, вентилация)"),
             ("табло", "т. 5 – главното електрическо табло"),
             ("инсталац", "т. 6 – пожароизвестителната и пожарогасителната инсталация"),
             ("опасност", "т. 7 – вероятните опасности"),
             ("връзка", "т. 8 – връзката с обекта"),
             ("водоизточн", "т. 10 – водоизточниците"),
             ("подстъп", "т. 11 – подстъпите"),
             ("звено", "т. 12 – обслужващото звено и времето за реагиране")]


def bad(msg):
    sys.stderr.write("BAD_INPUT: " + msg + "\n")
    sys.exit(2)


def load(base, path, what):
    if isinstance(path, dict):                         # inline – the plan JSON kept in the register is whole
        return json.loads(json.dumps(path))
    if not path:
        bad(f"липсва {what}")
    full = path if os.path.isabs(path) else os.path.join(base, path)
    try:
        return json.load(open(full, encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        bad(f"{what}: не мога да прочета {full}: {e}")


def num(v):
    if isinstance(v, float):
        v = round(v, 2)
        v = int(v) if v == int(v) else v
    return str(v).replace(".", ",")


def portrait(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0), Cm(2.0), Cm(1.5)
    return sec


def landscape(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page(sec)                                         # the margins of the sheets (left 2 cm for binding)
    return sec


def unpad(png):
    """The framed sheet without the white page margins pad_to_page added – for the page of the plan, whose own
    margins are the same."""
    from PIL import Image
    im = Image.open(png)
    dpi = (im.info.get("dpi") or (300, 300))[0]
    mm = dpi / 25.4
    l, r, t, b = (round(MARGINS[k] * 10 * mm) for k in ("left", "right", "top", "bottom"))
    out = os.path.splitext(png)[0] + "_list.png"
    im.crop((l, t, im.width - r, im.height - b)).save(out, dpi=(dpi, dpi))
    return out


def sheet(doc, png):
    """The graphic part – the whole framed sheet on a landscape page."""
    landscape(doc)
    path = unpad(png)
    p = B.fmt(doc.add_paragraph(), B.CENTER, line=None)
    p.add_run().add_picture(path, width=Cm(26.6))
    os.remove(path)


def back(doc, nomer, title, obekt):
    """The back of a sheet – the field for the number, the title of the table part, the object."""
    landscape(doc)
    number_box(doc, nomer)
    B.para(doc, title, align=B.CENTER, indent=False)
    if obekt:
        B.para(doc, obekt, align=B.CENTER, indent=False)
    B.blank(doc)


def pieces(n, one, few):
    return f"{num(n)} {one if n == 1 else few}"


def sili_sentence(R):
    """„Сили и средства.“ – from the result of the calculation."""
    s = R.get("struinik") or ""
    parts = [f"{pieces(R.get('struinitsi_za_gasene', 0), 'струйник', 'струйника')} „{s}“ за гасене"]
    if R.get("struinitsi_za_zashtita"):
        parts.append(f"{pieces(R['struinitsi_za_zashtita'], 'струйник', 'струйника')} „{s}“ за защита")
    parts.append(pieces(R.get("ekipi_obshto", 0), "екип", "екипа"))
    parts.append(f"разход на вода {num(R.get('razhod_obshto_l_s', R.get('razhod_za_gasene_l_s', 0)))} l/s")
    parts.append(pieces(R.get("pozharni_avtomobili", 0), "пожарен автомобил", "пожарни автомобила"))
    return "По разчета (Приложение № 5): " + ", ".join(parts) + "."


def items(doc, prefix, rows):
    for i, row in enumerate(rows, 1):
        if not (isinstance(row, (list, tuple)) and len(row) == 2):
            bad(f"{prefix}: всеки ред е [заглавие, текст]")
        B.para(doc, f"{prefix}{i}. *{row[0]}* {row[1]}")


def gaps_6_1(danni):
    heads = " ".join(str(r[0]).lower() for r in danni)
    return [f"данни за обекта: няма чл. 6, ал. 1, {what}" for key, what in DANNI_6_1 if key not in heads]


def build(plan, base, out_dir):
    if not isinstance(plan, dict):
        bad("очаква се JSON обект")
    varianti = plan.get("varianti") or []
    if not varianti:
        bad("varianti: поне един вариант (чл. 6, ал. 2) – конкретна схема и разчет")
    name = str(plan.get("fail") or "Plan_pozharogasene").strip()
    suffix = re.sub(r"^Plan_pozharogasene_?", "", name) or "obekt"
    year = plan.get("godina") or datetime.date.today().year
    obekt_line = ", ".join(x for x in (plan.get("obekt"), plan.get("adres")) if x)
    warnings, files = [], []

    # the sheets on their own – the same code the skill shemi uses
    sit = load(base, plan.get("situacionna"), "situacionna (ситуационната схема)")
    sit_out = os.path.join(out_dir, f"Situacionna_shema_{suffix}.docx")
    sit_res = situacionna.build(sit, sit_out)
    files += [sit_out, sit_res.get("dxf"), sit_res["grafichna"]]
    warnings += [f"ситуационна схема: {w}" for w in sit_res.get("preduprezhdeniya") or []]
    done = []
    for i, v in enumerate(varianti, 1):
        k = load(base, v.get("konkretna"), f"вариант {i}: konkretna")
        s = load(base, v.get("sili"), f"вариант {i}: sili")
        if not razchet_docx.valid(s):
            bad(f"вариант {i}: sili – очаква се JSON-ът на sili.py")
        fail = str(v.get("fail") or f"Konkretna_shema_{i}").strip()
        k_out = os.path.join(out_dir, fail + ".docx")
        k_res = konkretna_docx.build(k, k_out)
        r_out = os.path.join(out_dir, "Razchet_" + fail + ".docx")
        naslov = (k.get("list") or {}).get("naslov") or f"Конкретна схема № {i}"
        doc = B.new_document()                         # only the calculation – the back of the scheme's sheet
        razchet_docx.page(doc, s, naslov, obekt_line, v.get("pomeshtenie"), new_section=False)
        doc.save(r_out)
        B.report(r_out, letterhead=False)
        files += [k_out, k_res.get("dxf"), k_res["grafichna"], r_out]
        warnings += [f"вариант {i}: {w}" for w in k_res.get("preduprezhdeniya") or []]
        done.append((v, k, s, k_res, naslov))

    # the plan
    doc = B.new_document()
    B.letterhead(doc)
    B.approval(doc, plan.get("utvarzhdava") or B.director_approver(), year, reg=True)
    B.blank(doc)
    title = [f"на {plan.get('obekt')}," if plan.get("obekt") else ""] + \
            [x for x in (plan.get("adres"), plan.get("identifikator")) if x]
    B.title(doc, "П Л А Н  З А  П О Ж А Р О Г А С Е Н Е", [x for x in title if x])
    B.para(doc, "Разработен съгласно Указанията за разработване на планове за пожарогасене, утвърдени със заповед "
           "рег. № Iз-2749/18.11.2010 г. на главния секретар на МВР (чл. 5 – 9).")
    B.blank(doc, 2)
    sg = plan.get("saglasuvano")
    if sg:
        B.signature(doc, "", sg.get("lines"), sg.get("name"), left=[""], label="СЪГЛАСУВАНО:")
        B.blank(doc, 2)
    toc = ["1. Данни за обекта (чл. 6, ал. 1)",
           "2. Ситуационна схема – графична и таблична част (Приложения № 2а и № 2б)"]
    for i, (v, *_r) in enumerate(done, 1):
        toc.append(f"{2 + i}. Вариант № {i} – {v.get('ime') or ''}: решение, конкретна схема (Приложения № 3а и "
                   "№ 3б) и разчет на силите и средствата (Приложение № 5)")
    toc.append(f"{3 + len(done)}. Бланка за отчитане на произшествията и занятията (Приложение № 4)")
    B.para(doc, "Съдържание:", indent=False)
    for line in toc:
        B.para(doc, line, indent=False)

    portrait(doc)
    B.para(doc, "1. ДАННИ ЗА ОБЕКТА (чл. 6, ал. 1)", align=B.CENTER, indent=False)
    B.blank(doc)
    danni = plan.get("danni") or []
    items(doc, "1.", danni)
    warnings += gaps_6_1(danni)
    iz = plan.get("izgotvil") or {}
    B.blank(doc, 2)
    B.signature(doc, iz.get("data"), iz.get("lines"), iz.get("name"))
    if not iz.get("name"):
        warnings.append("няма име на изготвилия – оставено е място")

    lst = sit.get("list") or {}
    sheet(doc, sit_res["grafichna"])
    back(doc, lst.get("nomer"), "Ситуационна схема – таблична част (Приложение № 2б)", lst.get("obekt"))
    table_2b(doc, sit.get("tablitsa") or {})

    for i, (v, k, s, k_res, naslov) in enumerate(done, 1):
        n = 2 + i
        portrait(doc)
        B.para(doc, f"{n}. ВАРИАНТ № {i} – {str(v.get('ime') or '').upper()} (чл. 6, ал. 2)", align=B.CENTER,
               indent=False)
        B.blank(doc)
        rows = [list(r) for r in v.get("reshenie") or []]
        if not any("сили" in str(r[0]).lower() for r in rows if r):
            rows.insert(min(3, len(rows)), ["Сили и средства.", sili_sentence(s.get("rezultat") or {})])
        items(doc, f"{n}.", rows)
        if not v.get("reshenie"):
            warnings.append(f"вариант {i}: няма решение (мястото, развитието, схемата на подаване) – чл. 6, ал. 2")
        kl = k.get("list") or {}
        sheet(doc, k_res["grafichna"])
        back(doc, kl.get("nomer"), f"{naslov} – таблична част (Приложение № 3б)", kl.get("obekt"))
        table_3b(doc, k.get("tablitsa_3b") or {})
        razchet_docx.page(doc, s, naslov, plan.get("obekt"), v.get("pomeshtenie"))

    landscape(doc)
    B.para(doc, f"{3 + len(done)}. БЛАНКА ЗА ОТЧИТАНЕ НА ПРОИЗШЕСТВИЯТА И ЗАНЯТИЯТА (Приложение № 4 към чл. 7, ал. 5)",
           align=B.CENTER, indent=False)
    B.blank(doc)
    p4 = plan.get("prilozhenie_4") or {}
    table_4(doc, p4.get("redove"), int(p4.get("broi") or 14))
    B.page_numbers(doc)
    out = os.path.join(out_dir, name + ".docx")
    doc.save(out)
    B.report(out)
    return {"plan": out, "failove": [f for f in files if f], "preduprezhdeniya": list(dict.fromkeys(warnings))}


def main():
    ap = argparse.ArgumentParser(description="План за пожарогасене по Указанията Iз-2749/2010")
    ap.add_argument("plan", help="plan.json")
    ap.add_argument("-d", "--dir", required=True, help="папката за файловете")
    a = ap.parse_args()
    try:
        plan = json.load(open(a.plan, encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        bad(f"не мога да прочета {a.plan}: {e}")
    os.makedirs(a.dir, exist_ok=True)
    res = build(plan, os.path.dirname(os.path.abspath(a.plan)), a.dir)
    print(json.dumps(res, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
