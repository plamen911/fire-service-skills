#!/usr/bin/env python3
"""Smoke tests for build_ptz.py. Run from the skill folder:
    python3 tests/test_build.py
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
SCRIPT = os.path.join(SKILL, "scripts", "build_ptz.py")
FIXTURES = os.path.join(HERE, "fixtures")              # an invented site and scenario, not a real one

# the tests never touch the network: no key, and an invented staff list instead of the real one
os.environ["FIRE_SERVICE_NO_KEY"] = "1"
os.environ["FIRE_SERVICE_STAFF_CSV"] = os.path.abspath(os.path.join(SKILL, "..", "..", "..", "..", "tests", "fixtures", "staff.csv"))

from docx import Document  # noqa: E402

failures = []


def check(cond, msg):
    print(("  ok  " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


def run(data):
    tmp = tempfile.mkdtemp()
    shutil.copy(os.path.join(FIXTURES, "shema.json"), tmp)
    src = os.path.join(tmp, "data.json")
    with open(src, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    out = os.path.join(tmp, "out.docx")
    r = subprocess.run([sys.executable, SCRIPT, src, "-o", out], capture_output=True, text=True)
    return r, out


def text_of(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts), d


def fixture():
    return json.load(open(os.path.join(FIXTURES, "data.json"), encoding="utf-8"))


def test_full():
    print("build_ptz.py – the full plan with a scheme")
    r, out = run(fixture())
    check(r.returncode == 0, "builds without an error" + (": " + r.stderr[:300] if r.returncode else ""))
    if r.returncode:
        return
    check("DOC_CHECK" not in r.stderr, "the document keeps the common rules (no bold, letterhead, „–“, „№ “)"
          + (": " + r.stderr.strip()[:200] if "DOC_CHECK" in r.stderr else ""))
    res = json.loads(r.stdout)
    check(res["preduprezhdeniya"] == [], f"no warnings for a complete plan: {res['preduprezhdeniya']}")
    check(res["obshto_minuti"] == 77, f"the total time is the sum of the rows, 45″ rounded up ({res['obshto_minuti']})")
    text, doc = text_of(out)
    for want in ("УТВЪРЖДАВАМ:", "П Л А Н – К О Н С П Е К Т", "за провеждане на пожаро-тактическо занятие",
                 "със състава на II дежурна смяна при РСПБЗН – Левски",
                 "I. Тема: Гасене на пожар в склад за дървен материал", "Трениране на КЕ за ръководство на екипите",
                 "III. Метод: Практически.", "V. Време: от 09:00 ч. до 10:20 ч. на 28.10.2026 г.",
                 "VI. Участващи: личният състав на II ДС при РСПБЗН – Левски", "VII. Материално осигуряване: флагчета",
                 "ОЧАКВАНИ ДЕЙСТВИЯ", "Общо", "77′", "РАЗРАБОТИЛ:", "СХЕМА",
                 "за разполагане на силите и средствата при провеждане на ПТЗ в склад"):
        check(want in text, f"has „{want}“")
    check(len(doc.inline_shapes) == 1, "the scheme is in the document")
    check(res.get("dxf") == os.path.splitext(out)[0] + "_shema.dxf" and os.path.exists(res["dxf"]),
          "the scheme is also written as a DXF for AutoCAD / CorelDRAW beside the .docx")
    from docx.enum.section import WD_ORIENT
    check(len(doc.sections) == 2 and doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE
          and doc.sections[-1].page_width > doc.sections[-1].page_height, "the scheme is on a landscape page")
    check(doc.sections[0].orientation == WD_ORIENT.PORTRAIT, "the plan itself stays portrait")
    sys.path.insert(0, os.path.join(SKILL, "scripts"))
    import build_ptz
    check(build_ptz.orientation(110, 82) == "landscape", "a wide site – a landscape page")
    check(build_ptz.orientation(40, 90) == "portrait", "a tall site – a portrait page")
    check(build_ptz.orientation(40, 90, "landscape") == "landscape", "\"orientatsiya\" overrides the choice")
    for old in ("РСПАБ", "РОД", "диспечер", "бойна готовност", "ДБС", "полк."):
        check(old not in text, f"no old term „{old}“")


def test_old_terms_and_missing():
    print("build_ptz.py – old terms, no scheme, no address")
    d = fixture()
    for k in ("shema", "myasto", "approver"):
        d.pop(k)
    d["hod"][0]["deystviya"] = "Диспечерът приема съобщението за пожар и вдига ДБС по тревога."
    d["hod"][2]["deystviya"] = "РОД разпорежда на боец № 1 от I отделение да разузнае."
    r, out = run(d)
    check(r.returncode == 0, "builds")
    if r.returncode:
        return
    w = " | ".join(json.loads(r.stdout)["preduprezhdeniya"])
    for want in ("диспечер", "ДБС", "РОД", "отделени", "бо", "без схема", "адрес"):
        check(want.lower() in w.lower(), f"warns about „{want}“")
    text, _ = text_of(out)
    check("НАЧАЛНИК НА" in text and "РСПБЗН – ЛЕВСКИ:" in text, "the chief of the unit approves (from the staff list)")


def test_upbzn():
    print("build_ptz.py – УПБЗН")
    d = fixture()
    d["zveno"] = "УПБЗН – Сторгозия"
    d.pop("approver")
    d.pop("shema")
    r, out = run(d)
    check(r.returncode == 0, "builds for a УПБЗН")
    if r.returncode == 0:
        text, _ = text_of(out)
        check("УПБЗН – СТОРГОЗИЯ:" in text and "при УПБЗН – Сторгозия" in text, "the УПБЗН is in the approval and the title")


def test_kais_scheme():
    print("build_ptz.py – a scheme started from the cadastre (КАИС): contours, the parcel, street parcels")
    import ezdxf
    d = fixture()
    tmp = tempfile.mkdtemp()
    sh = {"teren": [120, 90], "kais": {"x0": 431100.0, "y0": 4808800.0, "imot": "99999.10.20", "data": "25.10.2026"},
          "imot": {"kontur": [[20, 20], [100, 20], [100, 70], [20, 70]], "etiket": "ПИ 99999.10.20"},
          "sgradi": [{"kontur": [[30, 30], [70, 30], [70, 50], [30, 50]], "etiket": "Склад", "etazhnost": "1 ет."}],
          "ulitsi": [{"kontur": [[0, 75], [120, 75], [120, 88], [0, 88]], "ime": "ул. „Примерна“", "ime_pos": [60, 81]}],
          "vodoiztochnitsi": [{"vid": "hidrant", "x": 10, "y": 80, "etiket": "ПХ 70/80"}],
          "pozhar": [{"x": 45, "y": 35, "w": 15, "h": 10}]}
    path = os.path.join(tmp, "kais_shema.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(sh, f, ensure_ascii=False)
    d["shema"] = path
    r, out = run(d)
    check(r.returncode == 0, "builds with a КАИС scheme" + (": " + r.stderr[:300] if r.returncode else ""))
    if r.returncode:
        return
    res = json.loads(r.stdout)
    check(res.get("dxf") and os.path.exists(res["dxf"]), "the DXF is written")
    if res.get("dxf"):
        doc = ezdxf.readfile(res["dxf"])
        xs = [p[0] for e in doc.modelspace().query("LWPOLYLINE") for p in e.get_points("xy")]
        check(xs and min(xs) > 431000, "the DXF is in КС2005 (the origin from \"kais\" is added back)")


def test_bad_input():
    print("build_ptz.py – bad input")
    for change, what in (({"zveno": ""}, "no unit"), ({"zveno": "РСПАБ – Левски"}, "an old unit name"),
                         ({"hod": []}, "no rows"), ({"hod": [{"obstanovka": "Пожар"}]}, "a row without actions"),
                         ({"shema": "missing.json"}, "a missing scheme file")):
        d = fixture()
        d.update(change)
        r, _ = run(d)
        check(r.returncode == 2 and "BAD_INPUT" in r.stderr, f"{what} → BAD_INPUT")


if __name__ == "__main__":
    test_full()
    test_old_terms_and_missing()
    test_upbzn()
    test_kais_scheme()
    test_bad_input()
    print("FAILURES: " + "; ".join(failures) if failures else "ALL OK")
    sys.exit(1 if failures else 0)
