#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_ptz.py — изгражда ПЛАН-КОНСПЕКТ за провеждане на ПОЖАРО-ТАКТИЧЕСКО ЗАНЯТИЕ (ПТЗ) с дежурна
смяна като .docx на бланката на РДПБЗН – Плевен: шапка, „УТВЪРЖДАВАМ“ (началникът на РСПБЗН или
УПБЗН), полета I–VII (тема, цел, метод, място, време, участващи, материално осигуряване), таблица
„Обстановка – време – очаквани действия“ с общото време, „РАЗРАБОТИЛ“ и на нова страница СХЕМА за
разполагане на силите и средствата (чертае я скилът shemi, цветна). Страницата на схемата се завърта
сама – легнала (landscape) за широк терен, изправена (portrait) за тесен или висок: така, че чертежът
да излезе най-голям; "orientatsiya": "landscape" / "portrait" я задава изрично.

    python3 scripts/build_ptz.py data.json -o /mnt/user-data/outputs/Plan-konspekt-PTZ_....docx

data.json (задължителни: "zveno", "obekt", "hod"; останалото има стойност по подразбиране):
{
  "zveno": "РСПБЗН – Левски",                 // или "УПБЗН – Сторгозия"; началникът му утвърждава
  "smyana": "I",                               // коя дежурна смяна
  "year": "2026",                              // по подразбиране – от датата
  "obekt": "склад за дървен материал на „Дървен свят“ ЕООД",   // „Гасене на пожар в …“
  "tema": "…",                                 // друга тема вместо „Гасене на пожар в {obekt}“
  "tsel": ["…", "…"],                          // по подразбиране – петте цели от образеца
  "metod": "Практически",
  "myasto": "гр. Левски, ул. „Примерна“ № 1",   // точният адрес на обекта
  "vreme": {"ot": "09:00", "do": "11:00", "data": "28.10.2026"},   // или готов текст
  "uchastvashti": "…",                         // по подразбиране – „личният състав на I ДС при … и представител на обекта“
  "materialno": "…",                           // по подразбиране – флагчетата от образеца, средства за имитация, план-конспект
  "hod": [                                     // редовете на таблицата, по реда на занятието
    {"obstanovka": "Пожар в склад за дървен материал", "vreme": "45″",
     "deystviya": "Дежурният в ОЦ приема съобщението за пожар, …"},
    {"obstanovka": "Движение до обекта", "vreme": 10, "deystviya": "…"}   // число – минути
  ],
  "obshto": "120′",                            // по подразбиране – сборът на времената
  "approver": {"lines": ["НАЧАЛНИК НА", "РСПБЗН – ЛЕВСКИ:", "ГЛАВЕН ИНСПЕКТОР"], "name": "Име Фамилия"},
                                               // по подразбиране – началникът на звеното от списъка на служителите
  "razrabotil": {"date": "20.10.2026 г.", "lines": ["ИНСПЕКТОР"], "name": "Име Фамилия"},
  "shema": "shema.json",                       // JSON за скила shemi (чертае се тук; началото му – от
                                               //   _shared/scripts/kais.py obekt --shema) или готов .png;
                                               //   пътят е спрямо папката на data.json; без него – без схема
  "shema_zaglavie": "…",                       // по подразбиране – „за разполагане на силите и средствата
                                               //   при провеждане на ПТЗ в {obekt}“
  "orientatsiya": "landscape"                  // страницата на схемата; по подразбиране – по пропорциите на терена
}

Времето в таблицата: число – минути („10′“); текст се пише както е („45″“, „15′ + 10′ движение до
РСПБЗН“). За общото време се събират всички стойности с ′ (минути) и ″ (секунди).

Типографията следва _shared/conventions.md: „№ 4“, тире „–“ и НИКАКЪВ получер шрифт. Скриптът
отпечатва JSON {"output", "redove", "obshto_minuti", "shema", "dxf", "preduprezhdeniya"}. Когато
схемата е JSON, до .docx се записват `…_shema.png` и същата схема като `…_shema.dxf` (AutoCAD/CorelDRAW;
при сгради от КАИС – в координати КС2005). Предупрежденията
казват за липсващи данни, за разминаване между общото време и часовете в „Време“, за стари термини
(РСПАБ, РОД, диспечер, бойна готовност, ДБС, отделение, боец, п.п. автомобил) и за предупрежденията
на схемата.

Грешки (на stderr, код ≠ 0):
    BAD_INPUT: ...   — липсва задължително поле, неверен ред в "hod", няма го файлът на схемата
    SHEMA_FAILED: ... — скилът shemi не можа да начертае схемата
"""
import argparse
import json
import math
import os
import re
import subprocess
import sys

from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

HERE = os.path.dirname(os.path.abspath(__file__))
SKILLS = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, os.path.join(SKILLS, "plan-konspekti", "scripts"))
import blanka as B  # noqa: E402  (the common blank of the training documents)

SHEMA_SCRIPT = os.path.join(SKILLS, "shemi", "scripts", "shema.py")
# The page of the scheme: (page width, page height, text width, height left for the picture under the title), cm;
# margins 1.5 cm left and 1.0 cm right, top and bottom – the scheme takes all the room the title leaves.
PAGES = {"landscape": (29.7, 21.0, 27.2, 17.6), "portrait": (21.0, 29.7, 18.5, 26.3)}

TSEL = ["Изучаване на оперативно-тактическите особености на обекта.",
        "Трениране на КЕ за ръководство на екипите.",
        "Проверка на взаимодействието между екипите в състава на дежурната смяна.",
        "Проверка на готовността на водачите на пожарните автомобили.",
        "Отработване на мерките за безопасност."]
MATERIALNO = ("флагчета (червено – за обозначаване на мястото и площта на пожара, черно – зона на задимяване, "
              "зелено – необходимост от евакуация, жълто – наличие на токсични вещества и продукти на горене), "
              "средства за имитация, план-конспект и др.")
# The old terms of the 2003 textbook and what is written today.
OLD_TERMS = [(r"\bРСПАБ\b|\bРУПБС\b|\bРСПБС\b", "РСПБЗН или УПБЗН"), (r"\bРОД\b", "РМ (ръководител на място)"),
             (r"(?i)\bдиспечер", "дежурният в ОЦ"), (r"(?i)\bбойн\w* готовност", "оперативната готовност"),
             (r"\bДБС\b", "ДС"), (r"(?i)\bотделени", "екип"), (r"\bКО\b", "КЕ (командир на екип)"),
             (r"(?i)\bполк\.", "без „полк.“"), (r"(?i)\bбо(ец|йц)", "пожарникар"),
             (r"(?i)\bп\.\s?п\.\s?автомобил", "пожарен автомобил (ПА)"), (r"\bППА\b", "ПА")]
TIME = re.compile(r"(\d+(?:[.,]\d+)?)\s*(′|'|’|мин\.?|min|″|\"|”|сек\.?|s\b)")


def bad(msg):
    sys.stderr.write("BAD_INPUT: " + msg + "\n")
    sys.exit(2)


def seconds(value):
    """The duration of a row in seconds, or None when it has no ′/″ value."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value) * 60
    found = TIME.findall(str(value or ""))
    if not found:
        return None
    total = 0.0
    for n, unit in found:
        n = float(n.replace(",", "."))
        total += n if unit in ("″", '"', "”", "сек", "сек.", "s") else n * 60
    return total


def show_time(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{value:g}′"
    return B.typo(str(value or "").strip())


def clock(t):
    """„9:00“, „09.00“ → (hours, minutes) or None."""
    m = re.match(r"^\s*(\d{1,2})[:.](\d{2})\s*$", str(t or ""))
    return (int(m.group(1)), int(m.group(2))) if m else None


def vreme_text(v, warnings):
    if isinstance(v, str) and v.strip():
        return B.typo(v.strip()), None
    v = v if isinstance(v, dict) else {}
    ot, do, data = (str(v.get(k) or "").strip() for k in ("ot", "do", "data"))
    if not (ot and do and data):
        warnings.append("„Време“ е непълно – подай \"vreme\": {\"ot\", \"do\", \"data\"}")
    text = f"от {ot or '……'} ч. до {do or '……'} ч. на {(data.rstrip(' г.') + ' г.') if data else '……………… г.'}"
    a, b = clock(ot), clock(do)
    span = ((b[0] * 60 + b[1]) - (a[0] * 60 + a[1])) if a and b else None
    return text, span


def unit_approver(unit, warnings):
    """The chief of the РСПБЗН (staff list, as in plan-konspekti) or of the УПБЗН (началник на участък)."""
    if unit.upper().startswith("РСПБЗН"):
        ap = B.unit_chief_approver(unit)
    else:
        ap = None
        mod = B.staff_module()
        if mod is not None:
            try:
                _, rows = mod.load()
                want = mod.n(unit)
                for r in rows:
                    if r["status"] == "заета" and mod.n(r["uchastak"]) == want and \
                            mod.n(r["zaemana_dlazhnost"] or r["dlazhnost_po_shtat"]) in ("началник на участък",
                                                                                          "впд началник на участък"):
                        vpd = mod.n(r["zaemana_dlazhnost"]).startswith("впд")
                        ap = {"lines": [("ВПД " if vpd else "") + "НАЧАЛНИК НА", unit.upper() + ":",
                                        r["zvanie"].split("/")[-1].strip().upper()], "name": r["ime_kratko"]}
                        break
            except BaseException:  # noqa: BLE001  (no key, no network)
                ap = None
    if not ap:
        warnings.append("няма данни за утвърждаващия (началника на звеното) – подай \"approver\"")
        ap = {"lines": ["НАЧАЛНИК НА", unit.upper() + ":", B.PLACEHOLDER_RANK], "name": B.PLACEHOLDER_NAME}
    return ap


def field(doc, roman, label, value):
    B.para(doc, f"{roman}. {label}: {value}", indent=False, keep=False)


def room(page, title=""):
    """(width, height) in cm left for the picture on the page of the scheme, under „СХЕМА“ and the title –
    a title that wraps onto more lines takes its height from the picture."""
    tw, ph = PAGES[page][2], PAGES[page][3]
    lines = max(1, math.ceil(len(title) / (tw * 4.4)))     # about 4.4 characters of 12 pt per cm
    return tw, ph - 0.3 - 0.6 * (lines - 1)


def orientation(width, height, forced=None, title=""):
    """The page on which a drawing of these proportions comes out bigger – landscape for a wide site,
    portrait for a narrow or tall one."""
    if forced in PAGES:
        return forced
    scale = {k: min(room(k, title)[0] / width, room(k, title)[1] / height) for k in PAGES}
    return "landscape" if scale["landscape"] >= scale["portrait"] else "portrait"


def run_shema(src, png, page, title=""):
    tw, h = room(page, title)
    r = subprocess.run([sys.executable, SHEMA_SCRIPT, src, "-o", png, "--shirina-cm", str(tw),
                        "--visochina-cm", str(h)], capture_output=True, text=True)
    if r.returncode != 0:
        sys.stderr.write("SHEMA_FAILED: " + (r.stderr.strip() or "без съобщение") + "\n")
        sys.exit(4)
    try:
        return json.loads(r.stdout)
    except json.JSONDecodeError:
        return {}


def draw_scheme(path, base, out, warnings, forced=None, title=""):
    """(the picture, the page orientation, the .dxf or None): a .png as it is, or a .json drawn by the skill shemi next to the
    output – first for a landscape page; when the site fits a portrait page better, again for it."""
    src = path if os.path.isabs(path) else os.path.join(base, path)
    if not os.path.exists(src):
        bad(f"няма го файлът на схемата: {src}")
    if not src.lower().endswith(".json"):
        try:
            from PIL import Image
            with Image.open(src) as im:
                w, h = im.size
        except Exception:  # noqa: BLE001
            w, h = 3, 2
        return src, orientation(w, h, forced, title), None
    png = os.path.splitext(out)[0] + "_shema.png"
    res = run_shema(src, png, "landscape", title)
    wu, hu = (res.get("elementi") or {}).get("razmer_m") or [3, 2]
    page = orientation(wu, hu, forced, title)
    if page != "landscape":
        res = run_shema(src, png, page, title)
    warnings += ["схема: " + w for w in res.get("preduprezhdeniya") or []]
    dxf = res.get("dxf")                               # the same scheme for AutoCAD / CorelDRAW, beside the .png
    if not dxf or not os.path.exists(dxf):
        warnings.append("схема: DXF не се записа – вижте скила shemi (ezdxf)")
        dxf = None
    return png, page, dxf


def old_terms(texts, warnings):
    seen = set()
    for t in texts:
        for pattern, now in OLD_TERMS:
            m = re.search(pattern, t)
            if m and pattern not in seen:
                seen.add(pattern)
                warnings.append(f"стар термин „{m.group(0)}“ – днес: {now}")


def build(d, out, base):
    warnings = []
    unit = B.typo(str(d.get("zveno") or "").strip())
    obekt = B.typo(str(d.get("obekt") or "").strip())
    hod = d.get("hod")
    if not unit:
        bad('липсва "zveno" (РСПБЗН – … или УПБЗН – …)')
    if not re.match(r"^(РСПБЗН|УПБЗН) – \S", unit):
        bad(f'"zveno" трябва да е „РСПБЗН – …“ или „УПБЗН – …“, а е „{unit}“')
    if not obekt and not d.get("tema"):
        bad('липсва "obekt" (или своя "tema")')
    if not isinstance(hod, list) or not hod:
        bad('липсва "hod" – редовете „обстановка – време – очаквани действия“')
    rows, total, untimed = [], 0.0, []
    for i, r in enumerate(hod, 1):
        if not isinstance(r, dict) or not str(r.get("obstanovka") or "").strip() or not str(r.get("deystviya") or "").strip():
            bad(f'hod[{i - 1}]: трябват "obstanovka" и "deystviya"')
        sec = seconds(r.get("vreme"))
        if sec is None:
            untimed.append(i)
        else:
            total += sec
        rows.append([f"{i}.", B.typo(str(r["obstanovka"]).strip()), show_time(r.get("vreme")),
                     B.typo(str(r["deystviya"]).strip())])
    if untimed:
        warnings.append("редове без време: " + ", ".join(map(str, untimed)))
    total_min = math.ceil(total / 60 - 1e-9)
    obshto = B.typo(str(d.get("obshto") or f"{total_min}′"))

    smyana = str(d.get("smyana") or "I").strip()
    vreme, span = vreme_text(d.get("vreme"), warnings)
    if span is not None and abs(span - total_min) > 10:
        warnings.append(f"„Време“ е {span} минути, а сборът в таблицата е {total_min} минути")
    year = str(d.get("year") or "").strip()
    if not year:
        m = re.search(r"(20\d\d)", json.dumps(d.get("vreme") or "", ensure_ascii=False))
        year = m.group(1) if m else "20……"
    myasto = B.typo(str(d.get("myasto") or "").strip())
    if not myasto:
        warnings.append("липсва точният адрес на обекта – подай \"myasto\"")
    approver = d.get("approver") or unit_approver(unit, warnings)
    sign = dict(d.get("razrabotil") or {})
    if not sign.get("name"):
        warnings.append("няма данни за разработилия – подай \"razrabotil\"")
    tema = B.typo(str(d.get("tema") or f"Гасене на пожар в {obekt}").strip().rstrip("."))
    tsel = [B.typo(str(t).strip()) for t in (d.get("tsel") or TSEL)]
    old_terms([tema, myasto, *tsel, str(d.get("uchastvashti") or ""), str(d.get("materialno") or "")] +
              [x for r in rows for x in r[1:]], warnings)

    doc = B.new_document()
    B.letterhead(doc, extra=unit)
    B.approval(doc, approver, year)
    B.title(doc, "ПЛАН–КОНСПЕКТ", ["за провеждане на пожаро-тактическо занятие",
                                   f"със състава на {smyana} дежурна смяна при {unit}"])
    field(doc, "I", "Тема", tema + ".")
    B.blank(doc)
    B.para(doc, "II. Цел:", indent=False)
    for i, t in enumerate(tsel, 1):
        B.para(doc, f"{i}. {t.rstrip('.')}.", indent=True)
    B.blank(doc)
    field(doc, "III", "Метод", B.typo(str(d.get("metod") or "Практически")).rstrip(".") + ".")
    field(doc, "IV", "Място", (myasto or "……………………………………………… (точният адрес на обекта)").rstrip(".") + ".")
    field(doc, "V", "Време", vreme.rstrip(".") + ".")
    uch = B.typo(str(d.get("uchastvashti") or f"личният състав на {smyana} ДС при {unit} и представител на обекта"))
    field(doc, "VI", "Участващи", uch.rstrip(".") + ".")
    field(doc, "VII", "Материално осигуряване", B.typo(str(d.get("materialno") or MATERIALNO)).rstrip(".") + ".")
    B.blank(doc)
    table = B.grid(doc, [["№", "по ред"], "ОБСТАНОВКА", "ВРЕМЕ", "ОЧАКВАНИ ДЕЙСТВИЯ"],
                   rows + [["", "Общо", obshto, ""]], [1.2, 4.2, 2.7, 8.4])
    for r, tr in enumerate(table.rows):
        if r == 0:
            continue
        for c in (0, 1, 2):
            for p in tr.cells[c].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    B.blank(doc, 2)
    B.signature(doc, sign.get("date") or "", sign.get("lines") or [], sign.get("name") or B.PLACEHOLDER_NAME,
                label="РАЗРАБОТИЛ:")

    shema, page, dxf = None, None, None
    if d.get("shema"):
        title2 = B.typo(str(d.get("shema_zaglavie") or
                            f"за разполагане на силите и средствата при провеждане на ПТЗ в {obekt or tema}"))
        shema, page, dxf = draw_scheme(str(d["shema"]), base, out, warnings, d.get("orientatsiya"), title2)
        pw, ph_ = PAGES[page][:2]
        text_w, pic_h = room(page, title2)
        sec = doc.add_section(WD_SECTION.NEW_PAGE)    # the scheme – on its own A4 page, turned to fit it
        sec.orientation = WD_ORIENT.LANDSCAPE if page == "landscape" else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(pw), Cm(ph_)
        sec.left_margin, sec.right_margin = Cm(1.5), Cm(1.0)
        sec.top_margin, sec.bottom_margin = Cm(1.0), Cm(1.0)
        sec.header_distance = sec.footer_distance = Cm(0.5)
        B.para(doc, "СХЕМА", align=B.CENTER, indent=False, size=14)
        tp = B.para(doc, title2, align=B.CENTER, indent=False)
        tp.paragraph_format.space_after = B.Pt(4)
        pic = B.fmt(doc.add_paragraph(), B.CENTER, line=None)
        width, height = Cm(text_w), None
        try:
            from PIL import Image
            with Image.open(shema) as im:
                w, h = im.size
            if h / w * text_w > pic_h:                 # limited by the height – keep it on the page
                width, height = None, Cm(pic_h)
        except Exception:  # noqa: BLE001
            pass
        pic.add_run().add_picture(shema, width=width, height=height)
    else:
        warnings.append("без схема – подай \"shema\" (JSON за скила shemi или готов .png)")

    os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
    doc.save(out)
    B.report(out)
    return {"output": out, "redove": len(rows), "obshto_minuti": total_min, "shema": shema, "dxf": dxf,
            "stranitsa_shema": page if shema else None, "preduprezhdeniya": warnings}


def main():
    ap = argparse.ArgumentParser(description="План-конспект за пожаро-тактическо занятие (ПТЗ)")
    ap.add_argument("data", help="data.json")
    ap.add_argument("-o", "--out", required=True, help="изходен .docx")
    a = ap.parse_args()
    try:
        with open(a.data, encoding="utf-8") as f:
            d = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        bad(f"не мога да прочета {a.data}: {e}")
    if not isinstance(d, dict):
        bad("data.json трябва да е JSON обект")
    res = build(d, a.out, os.path.dirname(os.path.abspath(a.data)))
    print(json.dumps(res, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
